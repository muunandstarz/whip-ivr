#!/usr/bin/env python3
"""Convert a Whip monthly utilization workbook into the LOU runtime dataset.

The workbook's reported Occupancy value is the canonical utilization value.  The
runtime dataset keeps both its raw counts and that reported percentage so the
application never derives a rate from an incompatible denominator.
"""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

VALID_CODES = {"RCK", "GB", "MIA", "ATL", "CHI", "ORL", "PHL", "BOS", "RVA", "DAL"}


def number(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    return float(str(value).replace(",", "").replace("%", "").strip())


def percent(value: Any) -> float:
    raw = number(value)
    if raw <= 1:
        raw *= 100
    if raw < 0 or raw > 100:
        raise ValueError(f"Reported occupancy is outside 0–100%: {value!r}")
    return round(raw, 1)


def iso_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=Path("server/data/utilization.json"))
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.workbook, data_only=True, read_only=True)
    worksheet = workbook["3_Utilizaton"]
    monthly: dict[str, list[dict[str, Any]]] = defaultdict(list)
    current_month: str | None = None
    max_input_rows = min(worksheet.max_row, 5_000)

    for row in worksheet.iter_rows(min_row=2, max_row=max_input_rows, max_col=11, values_only=True):
        review_period, code, period, market, available, rented, available_fleet, occupancy = row[:8]
        parsed_date = iso_date(review_period)
        if parsed_date:
            current_month = parsed_date
        if not isinstance(code, str) or code.strip() not in VALID_CODES or not current_month:
            continue

        fleet = number(available_fleet)
        rented_count = number(rented)
        available_count = number(available)
        reported_utilization = percent(occupancy)
        if fleet <= 0:
            if reported_utilization != 0:
                raise ValueError(f"{current_month} {code}: nonzero occupancy without an available fleet")
        else:
            calculated_utilization = round((rented_count / fleet) * 100, 1)
            # The source publishes occupancy rounded to whole percentage
            # points for most historical months.  Keep that reported value,
            # while refusing discrepancies larger than normal rounding.
            if abs(calculated_utilization - reported_utilization) > 0.51:
                raise ValueError(
                    f"{current_month} {code}: reported {reported_utilization}% does not match "
                    f"{rented_count}/{fleet} = {calculated_utilization}%"
                )

        monthly[current_month].append(
            {
                "code": code.strip(),
                "name": str(market or code).strip(),
                "period": str(period or code).strip(),
                "available": int(available_count),
                "rented": int(rented_count),
                "total": int(fleet),
                "utilization": reported_utilization,
            }
        )

    if not monthly:
        raise ValueError("No market utilization records found")

    month_data: dict[str, dict[str, Any]] = {}
    markets: dict[str, str] = {}
    for month_key in sorted(monthly):
        entries = sorted(monthly[month_key], key=lambda item: item["code"])
        # A source month must have a unique row per market.  This catches
        # malformed Excel ranges before they reach carrier-facing documents.
        codes = [entry["code"] for entry in entries]
        if len(codes) != len(set(codes)):
            raise ValueError(f"Duplicate market records in {month_key}")
        total_fleet = sum(entry["total"] for entry in entries)
        total_rented = sum(entry["rented"] for entry in entries)
        month_date = date.fromisoformat(month_key)
        month_data[month_key] = {
            "date": month_key,
            "month": month_date.strftime("%B %Y"),
            "year": month_date.year,
            "month_num": month_date.month,
            "available_fleet": sum(entry["available"] for entry in entries),
            "rented": total_rented,
            "total_available_fleet": total_fleet,
            "market_utilization": round((total_rented / total_fleet) * 100, 1) if total_fleet else 0,
            "markets": entries,
        }
        for entry in entries:
            markets[entry["code"]] = entry["name"]

    output = {
        "markets": dict(sorted(markets.items())),
        "monthly_data": month_data,
        "data_as_of": max(month_data),
        "source": {
            "file": args.workbook.name,
            "worksheet": worksheet.title,
            "reported_field": "Occupancy",
            "methodology": "The reported Occupancy field is retained as utilization. The denominator is the workbook's Available Fleet field, not the Available column.",
        },
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, indent=2) + "\n")
    total_rows = sum(len(value) for value in monthly.values())
    print(f"Imported {total_rows} market rows across {len(month_data)} report months through {output['data_as_of']}.")


if __name__ == "__main__":
    main()
